from openpyxl import load_workbook

def find_header_row(ws, expected_headers: list[str], max_scan_rows=20):
    expected = [h.strip().lower() for h in expected_headers]
    for r in range(1, max_scan_rows + 1):
        row_vals = [(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)]
        row_norm = [str(v).strip().lower() for v in row_vals]
        if all(h in row_norm for h in expected):
            return r, row_norm
    return None, None

def append_rows_to_sheet(template_path: str, output_path: str, sheet_name: str, rows: list[dict]):
    wb = load_workbook(template_path)
    ws = wb[sheet_name]

    # Adjust these to match your template headers exactly
    expected_headers = ["Raw table name", "Raw column name", "Data Type 2/ Precision", "Bronze table name", "Bronze column name"]
    header_row, header_norm = find_header_row(ws, expected_headers)

    if not header_row:
        raise ValueError(f"Could not find header row in sheet '{sheet_name}'")

    # Build column index map
    col_index = {}
    for idx, name in enumerate(header_norm, start=1):
        col_index[name] = idx

    def col_for(header_text: str):
        return col_index[header_text.strip().lower()]

    # Find first empty row after header
    r = header_row + 1
    while ws.cell(r, 1).value not in (None, ""):
        r += 1

    # Write rows
    for item in rows:
        if "data_entity" in item and "data entity" in col_index:
            ws.cell(r, col_for("data entity")).value = item.get("data_entity")
        if "source_field" in item and "source field" in col_index:
            ws.cell(r, col_for("source field")).value = item.get("source_field")

        if "raw_table" in item and "raw table name" in col_index:
            ws.cell(r, col_for("raw table name")).value = item.get("raw_table")
        if "raw_column" in item and "raw column name" in col_index:
            ws.cell(r, col_for("raw column name")).value = item.get("raw_column")

        if "bronze_table" in item and "bronze table name" in col_index:
            ws.cell(r, col_for("bronze table name")).value = item.get("bronze_table")
        if "bronze_column" in item and "bronze column name" in col_index:
            ws.cell(r, col_for("bronze column name")).value = item.get("bronze_column")

        r += 1

    wb.save(output_path)

